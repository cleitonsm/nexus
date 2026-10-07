import io
import unittest
from unittest.mock import patch

from src.domain import BlockType
from src.infrastructure.documents import extract_supported_document

try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover
    DocxDocument = None

_DOCX_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

_MARKDOWN = """# Politica de Ferias

Todo colaborador tem direito a ferias
apos doze meses de trabalho.

## Estagiarios

- Recesso de trinta dias.
- Recesso proporcional
  em contratos menores.

| Contrato | Recesso |
|----------|---------|
| 12 meses | 30 dias |
| 6 meses  | 15 dias |

```
# isto nao e titulo
linha de codigo
```

# Reembolso

O reembolso e feito em ate 10 dias uteis.
"""


def _extract_markdown(text: str):
    return extract_supported_document(
        filename="politica.md",
        content_type="text/markdown",
        raw_content=text.encode("utf-8"),
    )


class MarkdownExtractionTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.blocks = _extract_markdown(_MARKDOWN).blocks

    def test_headings_become_section_path_not_blocks(self) -> None:
        self.assertEqual(
            [block.type for block in self.blocks],
            [
                BlockType.PARAGRAPH,
                BlockType.LIST_ITEM,
                BlockType.LIST_ITEM,
                BlockType.TABLE,
                BlockType.CODE,
                BlockType.PARAGRAPH,
            ],
        )
        self.assertEqual(self.blocks[0].section_path, ("Politica de Ferias",))

    def test_wrapped_paragraph_lines_are_joined(self) -> None:
        self.assertEqual(
            self.blocks[0].text,
            "Todo colaborador tem direito a ferias apos doze meses de trabalho.",
        )

    def test_nested_heading_extends_section_path(self) -> None:
        self.assertEqual(
            self.blocks[1].section_path,
            ("Politica de Ferias", "Estagiarios"),
        )

    def test_list_item_keeps_continuation_line(self) -> None:
        self.assertEqual(
            self.blocks[2].text,
            "- Recesso proporcional em contratos menores.",
        )

    def test_table_keeps_header_lines_and_rows(self) -> None:
        table = self.blocks[3]
        self.assertEqual(table.header_lines, 2)
        self.assertEqual(len(table.rows), 2)
        self.assertIn("Contrato", table.header)

    def test_heading_marker_inside_code_fence_is_content(self) -> None:
        code = self.blocks[4]
        self.assertEqual(code.text, "# isto nao e titulo\nlinha de codigo")
        self.assertEqual(
            code.section_path,
            ("Politica de Ferias", "Estagiarios"),
        )

    def test_same_level_heading_replaces_deeper_sections(self) -> None:
        self.assertEqual(self.blocks[5].section_path, ("Reembolso",))

    def test_markdown_has_no_page(self) -> None:
        self.assertTrue(all(block.page is None for block in self.blocks))

    def test_text_before_first_heading_has_empty_section(self) -> None:
        blocks = _extract_markdown("Introducao solta.\n\n# Titulo\n\nCorpo.").blocks
        self.assertEqual(blocks[0].section_path, ())
        self.assertEqual(blocks[1].section_path, ("Titulo",))

    def test_skipped_heading_level_is_still_nested(self) -> None:
        blocks = _extract_markdown("# A\n\n### C\n\nCorpo.\n\n## B\n\nOutro.").blocks
        self.assertEqual(blocks[0].section_path, ("A", "C"))
        self.assertEqual(blocks[1].section_path, ("A", "B"))

    def test_table_without_separator_has_no_header(self) -> None:
        blocks = _extract_markdown("| a | b |\n| 1 | 2 |").blocks
        self.assertEqual(blocks[0].type, BlockType.TABLE)
        self.assertEqual(blocks[0].header_lines, 0)

    def test_unclosed_code_fence_is_kept(self) -> None:
        blocks = _extract_markdown("# A\n\n```\ncodigo sem fim").blocks
        self.assertEqual(blocks[0].type, BlockType.CODE)
        self.assertEqual(blocks[0].text, "codigo sem fim")


class PlainTextExtractionTestCase(unittest.TestCase):
    def test_paragraphs_are_split_on_blank_lines(self) -> None:
        document = extract_supported_document(
            filename="nota.txt",
            content_type="text/plain",
            raw_content=b"Primeiro paragrafo.\n\n\nSegundo paragrafo.\n",
        )
        self.assertEqual(
            [block.text for block in document.blocks],
            ["Primeiro paragrafo.", "Segundo paragrafo."],
        )
        self.assertEqual(document.blocks[0].section_path, ())

    def test_markdown_markers_are_not_interpreted_in_txt(self) -> None:
        document = extract_supported_document(
            filename="nota.txt",
            content_type="text/plain",
            raw_content=b"# nao e titulo\n\nCorpo.",
        )
        self.assertEqual(document.blocks[0].text, "# nao e titulo")
        self.assertEqual(document.blocks[1].section_path, ())


class PdfExtractionTestCase(unittest.TestCase):
    def _extract(self, pages: list[str]):
        with patch(
            "src.infrastructure.documents.structured_extraction._read_pdf_pages",
            return_value=pages,
        ):
            return extract_supported_document(
                filename="manual.pdf",
                content_type="application/pdf",
                raw_content=b"%PDF-fake",
            )

    def test_blocks_carry_one_based_page(self) -> None:
        document = self._extract(["Texto da primeira.", "Texto da segunda."])
        self.assertEqual([block.page for block in document.blocks], [1, 2])

    def test_numbered_headings_form_sections_across_pages(self) -> None:
        document = self._extract(
            [
                "1 Introducao\nO sistema atende areas.\n1.1 Escopo\nInclui chat.",
                "Continua o escopo aqui.\n2 Operacao\nRotina diaria.",
            ]
        )
        self.assertEqual(
            [(block.text, block.section_path, block.page) for block in document.blocks],
            [
                ("O sistema atende areas.", ("Introducao",), 1),
                ("Inclui chat.", ("Introducao", "Escopo"), 1),
                ("Continua o escopo aqui.", ("Introducao", "Escopo"), 2),
                ("Rotina diaria.", ("Operacao",), 2),
            ],
        )

    def test_numbered_sentence_is_not_a_heading(self) -> None:
        document = self._extract(["2 pessoas aprovaram o pedido.\nSegue o texto."])
        self.assertEqual(len(document.blocks), 1)
        self.assertEqual(document.blocks[0].section_path, ())

    def test_wrapped_lines_are_joined_and_blank_line_splits(self) -> None:
        document = self._extract(["Linha um\ncontinua.\n\nOutro paragrafo."])
        self.assertEqual(
            [block.text for block in document.blocks],
            ["Linha um continua.", "Outro paragrafo."],
        )

    def test_pdf_without_text_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            self._extract(["", "   "])


@unittest.skipIf(DocxDocument is None, "python-docx nao instalado")
class DocxExtractionTestCase(unittest.TestCase):
    def _extract(self):
        source = DocxDocument()
        source.add_heading("Politica de Ferias", level=1)
        source.add_paragraph("Direito apos doze meses.")
        source.add_heading("Estagiarios", level=2)
        source.add_paragraph("Recesso de trinta dias.", style="List Bullet")
        table = source.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Contrato"
        table.cell(0, 1).text = "Recesso"
        table.cell(1, 0).text = "12 meses"
        table.cell(1, 1).text = "30 dias"
        source.add_paragraph("Depois da tabela.")
        buffer = io.BytesIO()
        source.save(buffer)
        return extract_supported_document(
            filename="politica.docx",
            content_type=_DOCX_TYPE,
            raw_content=buffer.getvalue(),
        )

    def test_docx_preserves_order_sections_and_table(self) -> None:
        blocks = self._extract().blocks
        self.assertEqual(
            [block.type for block in blocks],
            [
                BlockType.PARAGRAPH,
                BlockType.LIST_ITEM,
                BlockType.TABLE,
                BlockType.PARAGRAPH,
            ],
        )
        self.assertEqual(blocks[0].section_path, ("Politica de Ferias",))
        self.assertEqual(
            blocks[1].section_path,
            ("Politica de Ferias", "Estagiarios"),
        )
        self.assertEqual(blocks[2].text, "Contrato | Recesso\n12 meses | 30 dias")
        self.assertEqual(blocks[2].header_lines, 1)
        self.assertEqual(blocks[3].text, "Depois da tabela.")


class ExtractionErrorsTestCase(unittest.TestCase):
    def test_unsupported_format_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            extract_supported_document(
                filename="dados.csv",
                content_type="text/csv",
                raw_content=b"a,b",
            )

    def test_empty_file_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            extract_supported_document(
                filename="vazio.md",
                content_type="text/markdown",
                raw_content=b"",
            )

    def test_file_with_only_headings_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            _extract_markdown("# Titulo\n\n## Outro\n")

    def test_non_utf8_text_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            extract_supported_document(
                filename="nota.txt",
                content_type="text/plain",
                raw_content=b"\xff\xfe\xfa",
            )


if __name__ == "__main__":
    unittest.main()
